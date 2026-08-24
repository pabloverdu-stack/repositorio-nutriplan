# -*- coding: utf-8 -*-
"""Calcula la nutricion de cada receta y exporta recetario.json + recetario.xlsx"""
import json, os
import nutricion as n
import ingredientes as I
from recetas_lote1 import RECETAS as DESAYUNOS
from recetas_comidas import RECETAS as COMIDAS
from recetas_meriendas import RECETAS as MERIENDAS
from recetas_lote2 import RECETAS as LOTE2
from recetas_lote3 import RECETAS as LOTE3
from generador import RECETAS as GENERADAS
from recetas_hamburguesas import RECETAS as HAMBURGUESAS
from recetas_elaboradas import RECETAS as ELABORADAS
from recetas_batidos import RECETAS as BATIDOS

BASE = os.path.dirname(os.path.abspath(__file__))
cat = n.cargar_catalogo()
_ING, _avisos = I.resolver(cat)
if _avisos:
    print("AVISO ingredientes sin resolver:", _avisos)

def _expandir(r):
    """Convierte recetas del lote2 (ingredientes por clave) al formato estandar."""
    if "_ings_clave" not in r:
        return r
    ings = []
    for clave, g, casera in r["_ings_clave"]:
        fid, _nom = _ING[clave]
        ings.append({"f_id": fid, "nombre": I.display(clave), "gramos": g, "casera": casera})
    r = {k: v for k, v in r.items() if k != "_ings_clave"}
    r["ingredientes"] = ings
    return r

RECETAS = [_expandir(r) for r in (DESAYUNOS + COMIDAS + MERIENDAS + HAMBURGUESAS + ELABORADAS + BATIDOS + LOTE2 + LOTE3 + GENERADAS)]

recetas_out = []
for r in RECETAS:
    ing = [{"f_id": i["f_id"], "g": i["gramos"]} for i in r["ingredientes"]]
    tot, faltan = n.calcular(cat, ing)
    # enriquecer ingredientes con el nombre oficial BEDCA
    ings = []
    for i in r["ingredientes"]:
        a = cat.get(i["f_id"], {})
        ings.append({**i, "bedca_nombre": a.get("nombre", "")})
    recetas_out.append({**r, "ingredientes": ings, "nutricion": tot, "nutrientes_incompletos": faltan})

# ---- JSON ----
out_json = os.path.join(BASE, "recetario.json")
json.dump({"fuente": "BEDCA", "n_recetas": len(recetas_out), "recetas": recetas_out},
          open(out_json, "w", encoding="utf-8"), ensure_ascii=False, indent=2)

# ---- Excel ----
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = Workbook()
# Hoja 1: resumen (una fila por receta con macros/micros)
ws = wb.active; ws.title = "Recetas"
cab = ["ID", "Nombre", "Tipo", "Dificultad", "Tiempo (min)", "Raciones"] + [lbl for _, lbl in n.NUTRIENTES]
ws.append(cab)
for r in recetas_out:
    fila = [r["id"], r["nombre"], r["tipo"], r["dificultad"], r["tiempo_min"], r["raciones"]]
    fila += [r["nutricion"][k] for k, _ in n.NUTRIENTES]
    ws.append(fila)

# Hoja 2: ingredientes (una fila por ingrediente)
ws2 = wb.create_sheet("Ingredientes")
ws2.append(["Receta ID", "Receta", "Ingrediente", "Gramos", "Medida casera", "f_id BEDCA", "Nombre BEDCA"])
for r in recetas_out:
    for i in r["ingredientes"]:
        ws2.append([r["id"], r["nombre"], i["nombre"], i["gramos"], i.get("casera", ""), i["f_id"], i["bedca_nombre"]])

# Hoja 3: elaboracion
ws3 = wb.create_sheet("Elaboracion")
ws3.append(["Receta ID", "Receta", "Paso", "Instruccion"])
for r in recetas_out:
    for idx, paso in enumerate(r["elaboracion"], 1):
        ws3.append([r["id"], r["nombre"], idx, paso])

# estilo cabeceras
head_fill = PatternFill("solid", fgColor="2E7D32")
head_font = Font(bold=True, color="FFFFFF")
thin = Side(style="thin", color="DDDDDD")
for sheet in (ws, ws2, ws3):
    for c in sheet[1]:
        c.fill = head_fill; c.font = head_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.freeze_panes = "A2"
    # anchos aproximados
    for col in sheet.columns:
        letter = col[0].column_letter
        maxlen = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        sheet.column_dimensions[letter].width = min(max(maxlen + 2, 10), 42)

out_xlsx = os.path.join(BASE, "recetario.xlsx")
try:
    wb.save(out_xlsx)
except PermissionError:
    out_xlsx = os.path.join(BASE, "recetario_nuevo.xlsx")
    wb.save(out_xlsx)
    print("(AVISO: recetario.xlsx estaba abierto; guardado como recetario_nuevo.xlsx)")
print("OK:", len(recetas_out), "recetas ->", out_json, "|", out_xlsx)
