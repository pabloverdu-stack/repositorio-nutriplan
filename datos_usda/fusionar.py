# -*- coding: utf-8 -*-
"""
Fusiona BEDCA + USDA y deja el resultado listo para el programa.

Qué hace:
 1. Lee el catálogo BEDCA y, para cada alimento con correspondencia revisada,
    rellena SOLO los huecos de cobre, manganeso, ác. pantoténico y azúcares.
 2. Escribe datos_usda/relleno.json, que es lo que consume nutricion.py.
 3. Recalcula la ficha nutricional de todas las recetas (recetario.json y
    app/data/recetario.json) sin tocar ninguna receta.
 4. Genera el informe para el usuario: informe_fusion_bedca_usda.xlsx

    python datos_usda/fusionar.py
"""
import json, io, os, sys, shutil, datetime

BASE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(BASE)
sys.path.insert(0, BASE)
sys.path.insert(0, REPO)

from correspondencias import MAPA, SIN_EMPAREJAR, NUTRIENTES_USDA

CLAVES = list(NUTRIENTES_USDA)
RELLENO = os.path.join(BASE, "relleno.json")
INFORME = os.path.join(REPO, "informe_fusion_bedca_usda.xlsx")

ETIQUETA = {"cobre_mg": "Cobre (mg)", "manganeso_mg": "Manganeso (mg)",
            "pantotenico_mg": "Ac. pantotenico B5 (mg)", "azucares_g": "Azucares (g)"}


def cargar():
    usda = json.load(io.open(os.path.join(BASE, "usda_sr_legacy.json"), encoding="utf-8"))
    import nutricion as n
    cat = n.cargar_catalogo(con_usda=False)   # BEDCA puro: los huecos se calculan sobre el original
    return usda, cat, n


def construir_relleno(usda, cat):
    """{f_id: {clave: valor}} solo con los huecos que se pueden rellenar."""
    relleno, filas, avisos = {}, [], []
    for fid, (fdc, tipo, excluir) in MAPA.items():
        if fid not in cat:
            avisos.append(f"f_id {fid} no existe en BEDCA"); continue
        u = usda.get(str(fdc))
        if not u:
            avisos.append(f"fdc_id {fdc} no existe en USDA (f_id {fid})"); continue
        bed = cat[fid]
        for k in CLAVES:
            if bed["nut"].get(k) is not None:      # BEDCA manda
                continue
            if k in excluir:
                continue
            v = u["nut"].get(k)
            if v is None:
                continue
            relleno.setdefault(fid, {})[k] = v
            filas.append([fid, bed["nombre"], bed.get("nombre_en", ""),
                          ETIQUETA[k], v, NUTRIENTES_USDA[k][2],
                          fdc, u["desc"],
                          "Equivalente" if tipo == "=" else "Aproximado",
                          "USDA SR Legacy 2018"])
    return relleno, filas, avisos


def _sumar(cat, ings, claves):
    """Como nutricion.calcular pero sobre la lista de claves que se le pase, para no
       perder claves antiguas de las recetas (p. ej. alcohol_g)."""
    tot = {k: 0.0 for k in claves}
    faltan = set()
    for ing in ings:
        a = cat[str(ing["f_id"])]
        factor = ing["g"] / 100.0
        for k in claves:
            if k in a["nut"]:
                tot[k] += a["nut"][k] * factor
            else:
                faltan.add(k)
    return {k: round(v, 2) for k, v in tot.items()}, sorted(faltan)


def recalcular_recetarios(cat_fusion):
    """Recalcula nutricion/nutrientes_incompletos in situ, sin perder recetas ni claves."""
    import nutricion as n
    resumen = []
    for ruta in (os.path.join(REPO, "recetario.json"),
                 os.path.join(REPO, "app", "data", "recetario.json")):
        if not os.path.exists(ruta):
            continue
        shutil.copy2(ruta, ruta + ".bak_prefusion")
        data = json.load(io.open(ruta, encoding="utf-8"))
        cambiadas = 0
        for r in data["recetas"]:
            ings = [{"f_id": i["f_id"], "g": i["gramos"]} for i in r.get("ingredientes", [])]
            if not ings:
                continue
            # se recalculan las claves estandar + las que ya tuviera la receta
            claves = list(dict.fromkeys(n.CLAVES + list(r.get("nutricion") or {})))
            try:
                tot, faltan = _sumar(cat_fusion, ings, claves)
            except KeyError:
                continue
            if r.get("nutricion") != tot:
                cambiadas += 1
            r["nutricion"] = tot
            r["nutrientes_incompletos"] = faltan
        data["fuente"] = "BEDCA + USDA SR Legacy (huecos de cobre, manganeso, B5 y azucares)"
        json.dump(data, io.open(ruta, "w", encoding="utf-8"), ensure_ascii=False, indent=2)
        resumen.append((ruta, len(data["recetas"]), cambiadas))
    return resumen


def escribir_informe(filas, cat, usda, avisos):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = Workbook()

    ws = wb.active; ws.title = "Valores anadidos"
    ws.append(["f_id BEDCA", "Alimento (BEDCA)", "Nombre en ingles (BEDCA)", "Nutriente",
               "Valor anadido", "Unidad", "fdc_id USDA", "Alimento equivalente (USDA)",
               "Tipo de equivalencia", "Fuente"])
    for f in sorted(filas, key=lambda x: (x[1], x[3])):
        ws.append(f)

    ws2 = wb.create_sheet("Resumen por alimento")
    ws2.append(["f_id BEDCA", "Alimento (BEDCA)", "fdc_id USDA", "Alimento equivalente (USDA)",
                "Tipo de equivalencia", "Nutrientes anadidos"])
    porali = {}
    for f in filas:
        porali.setdefault(f[0], [f[0], f[1], f[6], f[7], f[8], []])[5].append(f[3])
    for v in sorted(porali.values(), key=lambda x: x[1]):
        ws2.append(v[:5] + [", ".join(v[5])])

    ws3 = wb.create_sheet("Sin emparejar")
    ws3.append(["f_id BEDCA", "Alimento (BEDCA)", "Motivo"])
    for fid, motivo in SIN_EMPAREJAR.items():
        ws3.append([fid, cat.get(fid, {}).get("nombre", "?"), motivo])
    ws3.append(["", "Biotina (B8) en TODOS los alimentos",
                "Ningun banco de datos publico (BEDCA, USDA SR Legacy, CIQUAL) publica biotina "
                "de forma generalizada. Se queda sin rellenar y el programa la marca 's/d'."])
    for aviso in avisos:
        ws3.append(["", "", "AVISO: " + aviso])

    ws4 = wb.create_sheet("Metodo")
    metodo = [
        ["Informe de fusion BEDCA + USDA", ""],
        ["Generado", datetime.datetime.now().strftime("%d/%m/%Y %H:%M")],
        ["", ""],
        ["Por que", "BEDCA solo mide cobre en el 16 % de sus alimentos, manganeso en el 15 %, "
                    "acido pantotenico en el 13 % y azucares en el 17 %. Con esos huecos el "
                    "programa no podia decir si una dieta cubre esos nutrientes."],
        ["Fuente anadida", "USDA FoodData Central, SR Legacy (abril 2018). Cobertura: cobre 94 %, "
                           "manganeso 83 %, ac. pantotenico 82 %, azucares 77 %."],
        ["Regla", "Solo se rellenan HUECOS. Si BEDCA tiene el dato, se respeta BEDCA."],
        ["Alcance", "Los 187 alimentos que usa el recetario, emparejados y revisados uno a uno. "
                    "El resto del catalogo BEDCA se deja intacto."],
        ["Equivalencias", "'Equivalente' = mismo alimento y misma preparacion. "
                          "'Aproximado' = especie, corte o preparacion algo distinta "
                          "(p. ej. merluza -> whiting, dorada -> lubina cocinada)."],
        ["Biotina", "No se rellena: no existe fuente publica. El programa la marca 's/d'."],
        ["Reproducible", "python datos_usda/extraer_usda.py  y despues  python datos_usda/fusionar.py"],
        ["Correspondencias", "datos_usda/correspondencias.py (tabla editable, con un comentario por alimento)"],
    ]
    for fila in metodo:
        ws4.append(fila)

    cab = PatternFill("solid", fgColor="2E7D32")
    fnt = Font(bold=True, color="FFFFFF")
    for sh in (ws, ws2, ws3):
        for c in sh[1]:
            c.fill = cab; c.font = fnt
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sh.freeze_panes = "A2"
    for sh in wb.worksheets:
        for col in sh.columns:
            ancho = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            sh.column_dimensions[col[0].column_letter].width = min(max(ancho + 2, 12), 60)
    ws4.column_dimensions["B"].width = 90
    for fila in ws4.iter_rows():
        fila[-1].alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(INFORME)
    return INFORME


def main():
    usda, cat, n = cargar()
    relleno, filas, avisos = construir_relleno(usda, cat)
    json.dump(relleno, io.open(RELLENO, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"Relleno: {len(relleno)} alimentos, {len(filas)} valores -> {RELLENO}")
    for a in avisos:
        print("  AVISO:", a)

    # catálogo ya fusionado (nutricion.py aplica el relleno recién escrito)
    cat_fusion = n.cargar_catalogo(con_usda=True)
    for ruta in (os.path.join(REPO, "datos_bedca", "catalogo.json"),
                 os.path.join(REPO, "app", "data", "catalogo.json")):
        json.dump(cat_fusion, io.open(ruta, "w", encoding="utf-8"), ensure_ascii=False)
        print("Catalogo ->", ruta)

    for ruta, nrec, camb in recalcular_recetarios(cat_fusion):
        print(f"Recetario -> {ruta}  ({nrec} recetas, {camb} con nutricion actualizada)")

    print("Informe ->", escribir_informe(filas, cat, usda, avisos))


if __name__ == "__main__":
    main()
